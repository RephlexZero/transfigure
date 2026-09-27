#!/usr/bin/env -S uv run --script
# /// script
# requires-python = ">=3.11"
# dependencies = ["pillow>=11", "pillow-heif>=1", "openpyxl>=3.1"]
# ///
"""Regenerate the conversion test fixtures in this directory.

The fixtures are small files produced by *other* tools (ffmpeg, LibreOffice,
Pillow, libheif) so the converter is tested against real-world encoders
rather than only against its own output. They are committed; CI does not run
this script.

Requirements: uv, ffmpeg (with libmp3lame, libvorbis, flac, aac, alac and
x265-free HEIC support comes from pillow-heif), and LibreOffice (`soffice`,
with Writer and Calc). Python dependencies are declared inline and installed
by uv:

    uv run crates/converter/tests/fixtures/generate.py

Lint and type-check with `ruff check`, `ruff format` and `ty check`.
"""

import json
import os
import shutil
import subprocess
import tempfile

import openpyxl
import pillow_heif
from PIL import Image, ImageDraw

HERE = os.path.dirname(os.path.abspath(__file__))
MARKER = "Transfigure fixture — Café naïve 42"


def out(name):
    return os.path.join(HERE, name)


def write(name, text):
    with open(out(name), "w", encoding="utf-8", newline="") as f:
        f.write(text)


def run(*cmd):
    subprocess.run(
        cmd, check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL
    )


# ── Images ──────────────────────────────────────────────


def test_pattern(w=48, h=32, alpha=False):
    """Gradient with a solid block, so orientation and channels are checkable."""
    mode = "RGBA" if alpha else "RGB"
    img = Image.new(mode, (w, h))
    px = img.load()
    assert px is not None
    for y in range(h):
        for x in range(w):
            a = 255 if not alpha else (255 if x < w // 2 else 96)
            px[x, y] = (x * 255 // (w - 1), y * 255 // (h - 1), 128, a)[: len(mode)]
    ImageDraw.Draw(img).rectangle([4, 4, 11, 11], fill=(255, 0, 0, 255)[: len(mode)])
    return img


def images():
    rgb = test_pattern()
    rgba = test_pattern(alpha=True)

    rgba.save(out("image-rgba.png"))
    rgb.convert("L").save(out("image-gray.png"))
    rgb.convert("P", palette=Image.Palette.ADAPTIVE, colors=16).save(
        out("image-palette.png")
    )
    # 16-bit per channel PNG via ffmpeg (Pillow can't write 48-bit RGB).
    rgb.save(out("_tmp.png"))
    run(
        "ffmpeg",
        "-y",
        "-i",
        out("_tmp.png"),
        "-pix_fmt",
        "rgb48be",
        out("image-16bit.png"),
    )
    os.remove(out("_tmp.png"))

    rgb.save(out("image.jpg"), quality=90)
    rgb.convert("CMYK").save(out("image-cmyk.jpg"), quality=90)
    rgb.save(out("image-lossy.webp"), quality=80)
    rgba.save(out("image-lossless.webp"), lossless=True)
    rgb.save(out("image.bmp"))
    rgb.save(out("image.tiff"), compression="tiff_lzw")
    rgba.save(out("image.tga"))
    rgba.save(out("image.qoi"))
    rgba.save(out("image.ico"), sizes=[(16, 16), (32, 32), (48, 48)])
    rgba.save(out("image.dds"))
    rgba.save(out("image-bc1.dds"), pixel_format="DXT1")
    rgba.save(out("image-bc3.dds"), pixel_format="DXT5")

    frames = [test_pattern().rotate(i * 90, expand=False) for i in range(3)]
    frames[0].save(
        out("image-animated.gif"),
        save_all=True,
        append_images=frames[1:],
        duration=100,
        loop=0,
    )

    heif = pillow_heif.from_pillow(rgb)
    heif.save(out("image.heic"), quality=90)

    # OpenEXR and Radiance HDR via ffmpeg.
    rgb.save(out("_tmp.png"))
    run(
        "ffmpeg", "-y", "-i", out("_tmp.png"), "-pix_fmt", "gbrpf32le", out("image.exr")
    )
    run("ffmpeg", "-y", "-i", out("_tmp.png"), out("image.hdr"))
    os.remove(out("_tmp.png"))

    write(
        "image.svg",
        '<svg xmlns="http://www.w3.org/2000/svg" width="48" height="32" viewBox="0 0 48 32">'
        '<rect width="48" height="32" fill="#1c1917"/>'
        '<circle cx="16" cy="16" r="10" fill="#d98014"/>'
        '<path d="M30 6 42 26H18Z" fill="none" stroke="#fff" stroke-width="2"/>'
        "</svg>\n",
    )


# ── Audio ───────────────────────────────────────────────


def audio():
    # 0.6 s: 440 Hz left, 660 Hz right. Mono variant is 440 Hz.
    stereo = [
        "-f",
        "lavfi",
        "-i",
        "sine=f=440:d=0.6:r=44100",
        "-f",
        "lavfi",
        "-i",
        "sine=f=660:d=0.6:r=44100",
        "-filter_complex",
        "[0][1]amerge=inputs=2[a]",
        "-map",
        "[a]",
    ]
    mono = ["-f", "lavfi", "-i", "sine=f=440:d=0.6:r=44100"]
    jobs = [
        ("audio-stereo.mp3", stereo, ["-c:a", "libmp3lame", "-b:a", "96k"]),
        ("audio-mono.mp3", mono, ["-c:a", "libmp3lame", "-b:a", "64k"]),
        ("audio.flac", stereo, ["-c:a", "flac"]),
        ("audio.ogg", stereo, ["-c:a", "libvorbis", "-q:a", "2"]),
        ("audio-24bit.wav", stereo, ["-c:a", "pcm_s24le"]),
        ("audio-aac.m4a", stereo, ["-c:a", "aac", "-b:a", "96k"]),
        ("audio-alac.m4a", stereo, ["-c:a", "alac"]),
        ("audio.aac", stereo, ["-c:a", "aac", "-b:a", "96k", "-f", "adts"]),
        ("audio.aiff", stereo, ["-c:a", "pcm_s16be"]),
    ]
    for name, src, codec in jobs:
        run("ffmpeg", "-y", *src, *codec, out(name))


# ── Documents ───────────────────────────────────────────

MARKDOWN = f"""# {MARKER}

A paragraph with **bold**, *italic*, `code` and a [link](https://example.com).

## Lists

- First item
- Second item with ünïcödé
  - Nested item

1. One
2. Two

> A quotation.

| Name | Value |
| ---- | ----: |
| alpha | 1 |
| beta | 2 |

```
fn main() {{}}
```

---

Last paragraph.
"""

HTML = f"""<!DOCTYPE html>
<html><head><meta charset="utf-8"><title>Fixture</title>
<style>body {{ font-family: serif; }}</style>
<script>console.log("should not appear in text");</script></head>
<body>
<h1>{MARKER}</h1>
<p>A paragraph with <strong>bold</strong>, <em>italic</em>, <code>code</code> and a
<a href="https://example.com">link</a>. Ampersand &amp; entity &eacute;.</p>
<h2>Lists</h2>
<ul><li>First item</li><li>Second item</li></ul>
<ol><li>One</li><li>Two</li></ol>
<table><tr><th>Name</th><th>Value</th></tr><tr><td>alpha</td><td>1</td></tr></table>
</body></html>
"""


def documents():
    write("doc.md", MARKDOWN)
    write("doc.html", HTML)
    write("doc.txt", f"{MARKER}\n\nSecond paragraph of plain text.\nThird line.\n")

    tmp = tempfile.mkdtemp()
    try:
        src = os.path.join(tmp, "doc.html")
        with open(src, "w", encoding="utf-8") as f:
            f.write(HTML)
        for fmt, name in [
            ("docx:MS Word 2007 XML", "doc.docx"),
            ("odt", "doc.odt"),
            ("rtf:Rich Text Format", "doc.rtf"),
            ("pdf:writer_pdf_Export", "doc.pdf"),
        ]:
            run("soffice", "--headless", "--convert-to", fmt, "--outdir", tmp, src)
            ext = name.split(".")[1]
            shutil.copy(os.path.join(tmp, "doc." + ext), out(name))
    finally:
        shutil.rmtree(tmp)


# ── Data ────────────────────────────────────────────────

ROWS = [
    ["id", "name", "zip", "price", "note"],
    ["1", "Café", "00501", "9.99", "plain"],
    ["2", "Naïve, Inc.", "10001", "12", 'has "quotes"'],
    ["3", "Ωmega", "02134", "0.5", ""],
]


def data():
    import csv

    with open(out("data.csv"), "w", encoding="utf-8", newline="") as f:
        csv.writer(f).writerows(ROWS)
    with open(out("data.tsv"), "w", encoding="utf-8", newline="") as f:
        csv.writer(f, delimiter="\t", lineterminator="\n").writerows(ROWS)

    records: list[dict[str, object]] = [
        {
            "id": int(r[0]),
            "name": r[1],
            "zip": r[2],
            "price": float(r[3]),
            "note": r[4],
        }
        for r in ROWS[1:]
    ]
    with open(out("data.json"), "w", encoding="utf-8") as f:
        json.dump(records, f, ensure_ascii=False, indent=2)

    config = {
        "title": MARKER,
        "version": 3,
        "enabled": True,
        "ratio": 0.75,
        "tags": ["a", "b", "c"],
        "owner": {"name": "Tom", "email": "tom@example.com"},
        "servers": [{"host": "alpha", "port": 8001}, {"host": "beta", "port": 8002}],
        "empty_list": [],
        "tricky": "yes: no # not a comment",
        "numeric_string": "007",
    }
    with open(out("config.json"), "w", encoding="utf-8") as f:
        json.dump(config, f, ensure_ascii=False, indent=2)
    write(
        "config.yaml",
        f"""# comment
title: "{MARKER}"
version: 3
enabled: true
ratio: 0.75
tags: [a, b, c]
owner:
  name: Tom
  email: tom@example.com
servers:
  - host: alpha
    port: 8001
  - host: beta
    port: 8002
empty_list: []
tricky: "yes: no # not a comment"
numeric_string: "007"
""",
    )
    write(
        "config.toml",
        f"""# comment
title = "{MARKER}"
version = 3
enabled = true
ratio = 0.75
tags = ["a", "b", "c"]
empty_list = []
tricky = "yes: no # not a comment"
numeric_string = "007"

[owner]
name = "Tom"
email = "tom@example.com"

[[servers]]
host = "alpha"
port = 8001

[[servers]]
host = "beta"
port = 8002
""",
    )
    write(
        "config.xml",
        f"""<?xml version="1.0" encoding="UTF-8"?>
<config version="3">
  <title>{MARKER}</title>
  <enabled>true</enabled>
  <tags><tag>a</tag><tag>b</tag></tags>
  <owner name="Tom"><email>tom@example.com</email></owner>
</config>
""",
    )

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "People"
    for r in ROWS:
        ws.append(r)
    ws["D2"] = 9.99
    ws2 = wb.create_sheet("Second")
    ws2.append(["only", "sheet two"])
    wb.save(out("data.xlsx"))

    tmp = tempfile.mkdtemp()
    try:
        for fmt, ext in [("ods", "ods"), ("xls:MS Excel 97", "xls")]:
            run(
                "soffice",
                "--headless",
                "--convert-to",
                fmt,
                "--outdir",
                tmp,
                out("data.xlsx"),
            )
            shutil.copy(os.path.join(tmp, "data." + ext), out("data." + ext))
    finally:
        shutil.rmtree(tmp)

    import base64

    write("data.base64", base64.b64encode(MARKER.encode()).decode() + "\n")


if __name__ == "__main__":
    images()
    audio()
    documents()
    data()
    print("fixtures written to", HERE)
